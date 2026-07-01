"""
Workshido Master Template — v4
Design: matches professional worksheet image style
Header color + info bar + exercises A-E + tip box + footer
"""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\juand\Downloads\Workshido_Master_Template.docx"

# ── COLORS ────────────────────────────────────────────────────────────────────
C_HDR    = "1565C0"   # header blue
C_INFO   = "E3F2FD"   # info bar light blue
C_STUD   = "F5F5F5"   # student info bar gray
C_EX_A   = "1565C0"   # exercise A badge
C_EX_B   = "2E7D32"   # exercise B badge (green)
C_EX_C   = "6A1B9A"   # exercise C badge (purple)
C_EX_D   = "E65100"   # exercise D badge (orange)
C_EX_E   = "00695C"   # exercise E badge (teal)
C_TIP    = "FFF9C4"   # tip box yellow
C_THDR   = "1565C0"   # grammar table header
C_TALT   = "EBF5FB"   # grammar table alt row
C_WBANK  = "E8F5E9"   # word bank green
C_FOOTER = "1565C0"
WHITE    = "FFFFFF"
GRAY     = "757575"

# ── XML HELPERS ───────────────────────────────────────────────────────────────
def shd(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    tcPr.append(s)

def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr'); tbl.insert(0, tblPr)
    tb = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'), 'none')
        tb.append(b)
    tblPr.append(tb)

def cw(cell, cm):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    w = OxmlElement('w:tcW')
    w.set(qn('w:w'), str(int(cm * 567)))
    w.set(qn('w:type'), 'dxa')
    tcPr.append(w)

def rh(row, cm):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(int(cm * 567)))
    h.set(qn('w:hRule'), 'atLeast')
    trPr.append(h)

def va(cell, v='center'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    e = OxmlElement('w:vAlign'); e.set(qn('w:val'), v)
    tcPr.append(e)

def sp(para, before=0, after=0):
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after  = Pt(after)

def run(para, text, bold=False, italic=False, size=10, color=None):
    r = para.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(int(color[0:2],16),
                                    int(color[2:4],16),
                                    int(color[4:6],16))
    return r

def ph(para, text, size=10):
    """Placeholder: italic gray."""
    run(para, text, italic=True, size=size, color="AAAAAA")

def kill_spacing(doc):
    body = doc.element.body
    for el in body:
        tag = el.tag.split('}')[-1] if '}' in el.tag else el.tag
        if tag == 'p':
            pPr = el.find(qn('w:pPr'))
            if pPr is None:
                pPr = OxmlElement('w:pPr'); el.insert(0, pPr)
            spc = pPr.find(qn('w:spacing'))
            if spc is None:
                spc = OxmlElement('w:spacing'); pPr.append(spc)
            spc.set(qn('w:before'), '0'); spc.set(qn('w:after'), '0')

# ── EXERCISE LABEL HELPER ─────────────────────────────────────────────────────
def ex_label(doc, letter, color, instruction_placeholder):
    """Single-row table: colored letter badge + instruction text."""
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    badge = t.rows[0].cells[0]
    text  = t.rows[0].cells[1]
    cw(badge, 0.7); cw(text, 16.6)
    shd(badge, color); va(badge)
    p = badge.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 1, 1)
    run(p, letter, bold=True, size=12, color=WHITE)
    va(text)
    p2 = text.paragraphs[0]; sp(p2, 1, 1)
    run(p2, "  ", size=10)
    ph(p2, instruction_placeholder, size=10)
    return t

def ex_label_2col(doc, letter, color, instruction_placeholder):
    """Same but narrower for 2-column layouts."""
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    badge = t.rows[0].cells[0]
    text  = t.rows[0].cells[1]
    cw(badge, 0.65); cw(text, 8.35)
    shd(badge, color); va(badge)
    p = badge.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 1, 1)
    run(p, letter, bold=True, size=11, color=WHITE)
    va(text)
    p2 = text.paragraphs[0]; sp(p2, 1, 1)
    run(p2, "  ", size=10)
    ph(p2, instruction_placeholder, size=10)
    return t

# ═════════════════════════════════════════════════════════════════════════════
def create():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width    = Cm(21.0); sec.page_height   = Cm(29.7)
    sec.left_margin   = Cm(1.5);  sec.right_margin  = Cm(1.5)
    sec.top_margin    = Cm(1.5);  sec.bottom_margin = Cm(1.5)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10)

    # ── 1. HEADER BAR ─────────────────────────────────────────────────────────
    ht = doc.add_table(rows=1, cols=3)
    no_borders(ht)
    r0 = ht.rows[0]; rh(r0, 1.3)
    for c in r0.cells: shd(c, C_HDR)

    cw(r0.cells[0], 1.8); va(r0.cells[0])
    p = r0.cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p)
    ph(p, "{{NUMBER}}"); run(p, "", size=26, color=WHITE)
    # Redo — number placeholder white bold
    r0.cells[0].paragraphs[0].clear()
    p = r0.cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p)
    run(p, "{{NUMBER}}", bold=True, size=26, color=WHITE)

    cw(r0.cells[1], 12.8); va(r0.cells[1])
    p = r0.cells[1].paragraphs[0]; sp(p)
    run(p, "{{WORKSHEET_TITLE}}", bold=True, size=17, color=WHITE)

    cw(r0.cells[2], 2.7); va(r0.cells[2])
    p = r0.cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p)
    run(p, "{{LEVEL}}", bold=True, size=11, color=WHITE)

    # ── 2. INFO BAR (Goal / Skills / Time) ────────────────────────────────────
    it = doc.add_table(rows=1, cols=3)
    no_borders(it)
    r1 = it.rows[0]; rh(r1, 0.65)
    for c in r1.cells: shd(c, C_INFO)

    cw(r1.cells[0], 9.0); va(r1.cells[0])
    p = r1.cells[0].paragraphs[0]; sp(p, 2, 2)
    run(p, "Goal: ", bold=True, size=9)
    run(p, "{{GOAL}}", size=9)

    cw(r1.cells[1], 4.5); va(r1.cells[1])
    p = r1.cells[1].paragraphs[0]; sp(p, 2, 2)
    run(p, "Skills: ", bold=True, size=9)
    run(p, "{{SKILLS}}", size=9)

    cw(r1.cells[2], 3.8); va(r1.cells[2])
    p = r1.cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p, 2, 2)
    run(p, "Time: ", bold=True, size=9)
    run(p, "{{TIME}}", size=9)

    # ── 3. STUDENT INFO BAR (Name / Grade / Date) ─────────────────────────────
    st = doc.add_table(rows=1, cols=3)
    no_borders(st)
    r2 = st.rows[0]; rh(r2, 0.6)
    for c in r2.cells: shd(c, C_STUD)

    fields = [("Name:", 9.0), ("Grade:", 4.5), ("Date:", 3.8)]
    for i, (label, width) in enumerate(fields):
        cw(r2.cells[i], width); va(r2.cells[i])
        p = r2.cells[i].paragraphs[0]; sp(p, 2, 2)
        run(p, f"  {label} ", bold=True, size=9)
        run(p, "_" * (30 if i==0 else 16), size=9, color=GRAY)

    # Separator
    sep = doc.add_paragraph(); sp(sep, 0, 0)

    # ── 4. MAIN BODY: Left column (exercises A, B, C) + Right (grammar tip + D, E)
    body = doc.add_table(rows=1, cols=2)
    no_borders(body)
    left  = body.rows[0].cells[0]
    right = body.rows[0].cells[1]
    cw(left, 10.5); cw(right, 6.8)

    # Remove borders from body cells
    for cell in [left, right]:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tb = OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'), 'none')
            tb.append(b)
        tcPr.append(tb)

    # ── LEFT: Exercise A ─────────────────────────────────────────────────────
    p = left.paragraphs[0]; sp(p, 5, 2)
    run(p, "A", bold=True, size=13, color=C_EX_A)
    run(p, ".  ", bold=True, size=11)
    run(p, "{{EXERCISE_A_TITLE}}", bold=True, size=11)

    # Grammar conjugation table (template)
    gt = left.add_table(rows=5, cols=4)
    gt.style = 'Table Grid'
    gcws = [1.8, 2.8, 2.8, 3.1]
    hrow = gt.rows[0]; rh(hrow, 0.58)
    for i, (txt, w) in enumerate(zip(["", "{{COL1}}", "{{COL2}}", "{{COL3}}"], gcws)):
        shd(hrow.cells[i], C_THDR); cw(hrow.cells[i], w); va(hrow.cells[i])
        p = hrow.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 1, 1)
        run(p, txt, bold=True, size=8, color=WHITE)
    for di in range(4):
        drow = gt.rows[di+1]; rh(drow, 0.5)
        alt = C_TALT if di % 2 == 0 else WHITE
        for i, (txt, w) in enumerate(zip(
            [f"{{{{ROW{di+1}_SUBJ}}}}", f"{{{{ROW{di+1}_C1}}}}",
             f"{{{{ROW{di+1}_C2}}}}", f"{{{{ROW{di+1}_C3}}}}"], gcws)):
            shd(drow.cells[i], C_THDR if i==0 else alt)
            cw(drow.cells[i], w); va(drow.cells[i])
            p = drow.cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            sp(p, 1, 1)
            run(p, txt, bold=(i==0), size=8, color=WHITE if i==0 else None)

    # ── LEFT: Exercise B ─────────────────────────────────────────────────────
    p = left.add_paragraph(); sp(p, 7, 2)
    run(p, "B", bold=True, size=13, color=C_EX_B)
    run(p, ".  ", bold=True, size=11)
    run(p, "{{EXERCISE_B_TITLE}}", bold=True, size=11)

    for i in range(1, 7):
        p = left.add_paragraph(); sp(p, 1, 1)
        run(p, f"{i}. ", bold=True, size=10)
        run(p, f"{{{{EX_B_ITEM_{i}}}}}", size=10, color="AAAAAA")
        run(p, "  ( {{OPT1}} / {{OPT2}} )", italic=True, size=10, color="555555")

    # ── LEFT: Exercise C ─────────────────────────────────────────────────────
    p = left.add_paragraph(); sp(p, 7, 2)
    run(p, "C", bold=True, size=13, color=C_EX_C)
    run(p, ".  ", bold=True, size=11)
    run(p, "{{EXERCISE_C_TITLE}}", bold=True, size=11)

    for i in range(1, 6):
        p = left.add_paragraph(); sp(p, 2, 1)
        run(p, f"{i}. ", bold=True, size=10)
        run(p, f"{{{{EX_C_ITEM_{i}}}}}  ", size=10, color="AAAAAA")
        run(p, "______________________________", size=10)

    # ── RIGHT: Grammar Tip Box ────────────────────────────────────────────────
    p = right.paragraphs[0]; sp(p, 5, 0)
    run(p, "", size=5)  # spacer

    tip_tbl = right.add_table(rows=1, cols=1)
    tip_tbl.style = 'Table Grid'
    tc = tip_tbl.rows[0].cells[0]
    shd(tc, C_TIP)

    p = tc.paragraphs[0]; sp(p, 4, 2)
    run(p, "Remember!", bold=True, size=10, color=C_HDR)

    tip_lines = [
        ("He / She / It", " → verb + -s / -es"),
        ("(+)", "  {{POSITIVE_FORM}}"),
        ("(−)", "  {{NEGATIVE_FORM}}"),
        ("(?)", "  {{QUESTION_FORM}}"),
        ("Signals:", "  {{SIGNAL_WORDS}}"),
    ]
    for label, text in tip_lines:
        p = tc.add_paragraph(); sp(p, 2, 1)
        run(p, label, bold=True, size=8.5, color=C_HDR)
        run(p, text, size=8.5)

    # ── RIGHT: Exercise D ─────────────────────────────────────────────────────
    p = right.add_paragraph(); sp(p, 8, 2)
    run(p, "D", bold=True, size=13, color=C_EX_D)
    run(p, ".  ", bold=True, size=11)
    run(p, "{{EXERCISE_D_TITLE}}", bold=True, size=11)

    for i in range(1, 6):
        p = right.add_paragraph(); sp(p, 2, 1)
        run(p, f"{i}. ", bold=True, size=10)
        run(p, f"{{{{EX_D_ITEM_{i}}}}}  ", size=10, color="AAAAAA")
        run(p, "_______________________", size=10)

    # ── RIGHT: Exercise E ─────────────────────────────────────────────────────
    p = right.add_paragraph(); sp(p, 7, 2)
    run(p, "E", bold=True, size=13, color=C_EX_E)
    run(p, ".  ", bold=True, size=11)
    run(p, "{{EXERCISE_E_TITLE}}", bold=True, size=11)

    for i in range(1, 5):
        p = right.add_paragraph(); sp(p, 2, 1)
        run(p, f"{i}. ", bold=True, size=10)
        run(p, f"{{{{EX_E_ITEM_{i}}}}}?   ", size=10, color="AAAAAA")
        run(p, f"{{{{EX_E_ANS_{i}}}}}", bold=True, size=10, color=C_EX_A)

    # ── 5. WORD BANK ──────────────────────────────────────────────────────────
    p = doc.add_paragraph(); sp(p, 6, 2)
    run(p, "Word Bank:", bold=True, size=10, color=C_EX_B)

    wt = doc.add_table(rows=1, cols=1)
    wt.style = 'Table Grid'
    shd(wt.rows[0].cells[0], C_WBANK)
    p = wt.rows[0].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 4, 4)
    run(p, "{{WORD1}}  ◆  {{WORD2}}  ◆  {{WORD3}}  ◆  {{WORD4}}  ◆  {{WORD5}}  ◆  {{WORD6}}  ◆  {{WORD7}}  ◆  {{WORD8}}", size=10)

    # ── 6. FOOTER ─────────────────────────────────────────────────────────────
    ft = doc.add_table(rows=1, cols=3)
    no_borders(ft)
    fr = ft.rows[0]; rh(fr, 0.5)
    for c in fr.cells: shd(c, C_FOOTER)

    cw(fr.cells[0], 7.0); va(fr.cells[0])
    p = fr.cells[0].paragraphs[0]; sp(p, 2, 2)
    run(p, "  {{MOTIVATIONAL_NOTE}}", italic=True, size=8, color=WHITE)

    cw(fr.cells[1], 5.5); va(fr.cells[1])
    p = fr.cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 2, 2)
    run(p, "workshido.com", bold=True, size=9, color=WHITE)

    cw(fr.cells[2], 4.8); va(fr.cells[2])
    p = fr.cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p, 2, 2)
    run(p, "{{ANSWER_KEY_NOTE}}  ", size=8, color=WHITE)

    # ── Kill phantom spacing ───────────────────────────────────────────────────
    kill_spacing(doc)

    doc.save(OUT)
    print(f"Saved: {OUT}")

create()
