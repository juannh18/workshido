"""
Workshido – Professional Worksheet Generator
Page 1: header + exercises A-E + word bank + footer
Page 2: Answer Key
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = r"C:\Users\juand\Downloads\PresentSimple_Grammar_A1A2.docx"

C_HEADER = "1565C0"
C_INFO   = "E3F2FD"
C_STUD   = "F5F5F5"
C_THDR   = "1565C0"
C_TALT   = "EBF5FB"
C_EX_A   = "1565C0"
C_EX_B   = "2E7D32"
C_EX_C   = "6A1B9A"
C_EX_D   = "E65100"
C_EX_E   = "00695C"
C_WBANK  = "FFF9C4"
C_FOOTER = "1565C0"
WHITE    = "FFFFFF"
GRAY     = "888888"

# ── XML HELPERS ───────────────────────────────────────────────────────────────
def shd(cell, fill):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    tcPr.append(s)

def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr'); tbl.insert(0, tblPr)
    tb = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'), 'none'); tb.append(b)
    tblPr.append(tb)

def cw(cell, cm):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    w = OxmlElement('w:tcW')
    w.set(qn('w:w'), str(int(cm * 567))); w.set(qn('w:type'), 'dxa')
    tcPr.append(w)

def rh(row, cm, rule='atLeast'):
    tr = row._tr; trPr = tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(int(cm * 567))); h.set(qn('w:hRule'), rule)
    trPr.append(h)

def va(cell, v='center'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    e = OxmlElement('w:vAlign'); e.set(qn('w:val'), v); tcPr.append(e)

def sp(para, before=0, after=0):
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after  = Pt(after)

def r(para, text, bold=False, italic=False, size=10, color=None):
    x = para.add_run(text)
    x.bold = bold; x.italic = italic; x.font.size = Pt(size)
    if color:
        x.font.color.rgb = RGBColor(int(color[0:2],16), int(color[2:4],16), int(color[4:6],16))
    return x

def kill_spacing(doc):
    body = doc.element.body
    for el in body:
        tag = el.tag.split('}')[-1] if '}' in el.tag else el.tag
        if tag == 'p':
            pPr = el.find(qn('w:pPr'))
            if pPr is None:
                pPr = OxmlElement('w:pPr'); el.insert(0, pPr)
            spc = pPr.find(qn('w:spacing'))
            if spc is None:
                spc = OxmlElement('w:spacing'); pPr.append(spc)
            spc.set(qn('w:before'), '0'); spc.set(qn('w:after'), '0')

# ── DOCUMENT SETUP ────────────────────────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_width    = Cm(21.0); sec.page_height   = Cm(29.7)
sec.left_margin   = Cm(1.5);  sec.right_margin  = Cm(1.5)
sec.top_margin    = Cm(1.5);  sec.bottom_margin = Cm(1.5)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)

# ══════════════════════════════════════════════════════════════════════════════
# PAGE 1
# ══════════════════════════════════════════════════════════════════════════════

# ── HEADER BAR ────────────────────────────────────────────────────────────────
ht = doc.add_table(rows=1, cols=3); no_borders(ht)
r0 = ht.rows[0]; rh(r0, 1.2, 'exact')
for c in r0.cells: shd(c, C_HEADER)

cw(r0.cells[0], 1.8); va(r0.cells[0])
p = r0.cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p)
r(p, "1", bold=True, size=26, color=WHITE)

cw(r0.cells[1], 12.8); va(r0.cells[1])
p = r0.cells[1].paragraphs[0]; sp(p)
r(p, "PRESENT SIMPLE – GRAMMAR", bold=True, size=16, color=WHITE)

cw(r0.cells[2], 2.7); va(r0.cells[2])
p = r0.cells[2].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p)
r(p, "A1–A2", bold=True, size=11, color=WHITE)

# ── INFO BAR ──────────────────────────────────────────────────────────────────
it = doc.add_table(rows=1, cols=3); no_borders(it)
r1 = it.rows[0]; rh(r1, 0.6, 'exact')
for c in r1.cells: shd(c, C_INFO)

cw(r1.cells[0], 9.5); va(r1.cells[0])
p = r1.cells[0].paragraphs[0]; sp(p, 1, 1)
r(p, "Goal: ", bold=True, size=9)
r(p, "to use the Present Simple correctly in affirmative, negative and questions.", size=9)

cw(r1.cells[1], 3.8); va(r1.cells[1])
p = r1.cells[1].paragraphs[0]; sp(p, 1, 1)
r(p, "Skills: ", bold=True, size=9); r(p, "Grammar", size=9)

cw(r1.cells[2], 4.0); va(r1.cells[2])
p = r1.cells[2].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p, 1, 1)
r(p, "Time: ", bold=True, size=9); r(p, "60 minutes", size=9)

# ── STUDENT INFO BAR ──────────────────────────────────────────────────────────
si = doc.add_table(rows=1, cols=3); no_borders(si)
r2 = si.rows[0]; rh(r2, 0.55, 'exact')
for c in r2.cells: shd(c, C_STUD)

cw(r2.cells[0], 9.5); va(r2.cells[0])
p = r2.cells[0].paragraphs[0]; sp(p, 1, 1)
r(p, "  Name: ", bold=True, size=9); r(p, "_" * 35, size=9, color=GRAY)

cw(r2.cells[1], 3.8); va(r2.cells[1])
p = r2.cells[1].paragraphs[0]; sp(p, 1, 1)
r(p, "  Grade: ", bold=True, size=9); r(p, "_" * 12, size=9, color=GRAY)

cw(r2.cells[2], 4.0); va(r2.cells[2])
p = r2.cells[2].paragraphs[0]; sp(p, 1, 1)
r(p, "  Date: ", bold=True, size=9); r(p, "_" * 15, size=9, color=GRAY)

# ── EXERCISE A: Grammar table ─────────────────────────────────────────────────
p = doc.add_paragraph(); sp(p, 5, 2)
r(p, "A.  ", bold=True, size=11, color=C_EX_A)
r(p, "Complete the table.", bold=True, size=11)

gt = doc.add_table(rows=8, cols=4); gt.style = 'Table Grid'
gcws = [1.8, 4.9, 5.0, 5.6]

hrow = gt.rows[0]; rh(hrow, 0.56, 'exact')
for i, (txt, w) in enumerate(zip(["", "AFFIRMATIVE", "NEGATIVE", "QUESTION"], gcws)):
    shd(hrow.cells[i], C_THDR); cw(hrow.cells[i], w); va(hrow.cells[i])
    p = hrow.cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 1, 1)
    r(p, txt, bold=True, size=8.5, color=WHITE)

data = [
    ("I",    "I play.",    "I __________ play.",    "__________ I play?"),
    ("You",  "You play.",  "You __________ play.",  "__________ you play?"),
    ("He",   "He plays.",  "He __________ play.",   "__________ he play?"),
    ("She",  "She plays.", "She __________ play.",  "__________ she play?"),
    ("It",   "It plays.",  "It __________ play.",   "__________ it play?"),
    ("We",   "We play.",   "We __________ play.",   "__________ we play?"),
    ("They", "They play.", "They __________ play.", "__________ they play?"),
]
for di, (subj, aff, neg, q) in enumerate(data):
    drow = gt.rows[di+1]; rh(drow, 0.48, 'exact')
    alt = C_TALT if di % 2 == 0 else WHITE
    for i, (txt, w) in enumerate(zip([subj, aff, neg, q], gcws)):
        shd(drow.cells[i], C_THDR if i==0 else alt)
        cw(drow.cells[i], w); va(drow.cells[i])
        p = drow.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        sp(p, 1, 1)
        r(p, txt, bold=(i==0), size=8.5, color=WHITE if i==0 else None)

# ── EXERCISE B + C side by side ───────────────────────────────────────────────
bc = doc.add_table(rows=1, cols=2); no_borders(bc)
left_bc = bc.rows[0].cells[0]; right_bc = bc.rows[0].cells[1]
cw(left_bc, 9.5); cw(right_bc, 7.8)
for cell in [left_bc, right_bc]:
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tb = OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'), 'none'); tb.append(b)
    tcPr.append(tb)

# B left
p = left_bc.paragraphs[0]; sp(p, 6, 2)
r(p, "B.  ", bold=True, size=11, color=C_EX_B)
r(p, "Circle the correct option.", bold=True, size=11)

b_items = [
    "He ( play / plays ) football every Saturday.",
    "I ( don't / doesn't ) like coffee.",
    "( Do / Does ) she study English?",
    "They ( doesn't / don't ) go to school on Sundays.",
    "My sister ( watch / watches ) TV in the evening.",
    "We ( do / does ) our homework after dinner.",
]
for i, item in enumerate(b_items):
    p = left_bc.add_paragraph(); sp(p, 1, 1)
    r(p, f"{i+1}. ", bold=True, size=9.5)
    r(p, item, size=9.5)

# C right
p = right_bc.paragraphs[0]; sp(p, 6, 2)
r(p, "C.  ", bold=True, size=11, color=C_EX_C)
r(p, "Fill in the blanks.", bold=True, size=11)

c_items = [
    "I __________ (get up) at 7 every day.",
    "She __________ (not/like) spicy food.",
    "They __________ (go) to the gym.",
    "He __________ (have) a shower.",
    "We __________ (watch) a movie Fridays.",
]
for i, item in enumerate(c_items):
    p = right_bc.add_paragraph(); sp(p, 2, 2)
    r(p, f"{i+1}. ", bold=True, size=9.5)
    r(p, item, size=9.5)

# ── EXERCISE D + E side by side ───────────────────────────────────────────────
de = doc.add_table(rows=1, cols=2); no_borders(de)
left_de = de.rows[0].cells[0]; right_de = de.rows[0].cells[1]
cw(left_de, 9.5); cw(right_de, 7.8)
for cell in [left_de, right_de]:
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tb = OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'), 'none'); tb.append(b)
    tcPr.append(tb)

# D left
p = left_de.paragraphs[0]; sp(p, 6, 2)
r(p, "D.  ", bold=True, size=11, color=C_EX_D)
r(p, "Correct the mistakes.", bold=True, size=11)

d_items = [
    "She don't like ice cream.",
    "Do he play tennis?",
    "They doesn't have a car.",
    "I goes to bed at 10 pm.",
    "My brother play the guitar.",
]
for i, item in enumerate(d_items):
    p = left_de.add_paragraph(); sp(p, 2, 1)
    r(p, f"{i+1}. ", bold=True, size=9.5)
    r(p, item + "   ", italic=True, size=9.5)
    r(p, "________________________", size=9.5)

# E right
p = right_de.paragraphs[0]; sp(p, 6, 2)
r(p, "E.  ", bold=True, size=11, color=C_EX_E)
r(p, "Write the questions.", bold=True, size=11)

e_answers = ["Yes, I do.", "No, she doesn't.", "Yes, they do.", "No, he doesn't."]
for i, ans in enumerate(e_answers):
    p = right_de.add_paragraph(); sp(p, 2, 1)
    r(p, f"{i+1}. ", bold=True, size=9.5)
    r(p, "____________________________?  ", size=9.5)
    r(p, ans, bold=True, size=9.5, color=C_EX_A)

# ── WORD BANK ─────────────────────────────────────────────────────────────────
p = doc.add_paragraph(); sp(p, 6, 2)
r(p, "Word Bank:", bold=True, size=9.5, color=C_EX_B)

wt = doc.add_table(rows=1, cols=1); wt.style = 'Table Grid'
shd(wt.rows[0].cells[0], C_WBANK)
p = wt.rows[0].cells[0].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 3, 3)
words = ["do", "does", "don't", "doesn't", "plays", "watches", "goes", "has", "studies", "likes"]
r(p, "   ◆   ".join(words), size=9.5)

# ── FOOTER (page 1) ───────────────────────────────────────────────────────────
ft = doc.add_table(rows=1, cols=3); no_borders(ft)
fr = ft.rows[0]; rh(fr, 0.48, 'exact')
for c in fr.cells: shd(c, C_FOOTER)

cw(fr.cells[0], 7.0); va(fr.cells[0])
p = fr.cells[0].paragraphs[0]; sp(p, 2, 2)
r(p, "  Answer Key on page 2", italic=True, size=8, color=WHITE)

cw(fr.cells[1], 5.5); va(fr.cells[1])
p = fr.cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 2, 2)
r(p, "workshido.com", bold=True, size=9, color=WHITE)

cw(fr.cells[2], 4.8); va(fr.cells[2])
p = fr.cells[2].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p, 2, 2)
r(p, "Page 1  ", size=8, color=WHITE)

# ══════════════════════════════════════════════════════════════════════════════
# PAGE 2 — ANSWER KEY
# ══════════════════════════════════════════════════════════════════════════════
pb = doc.add_paragraph(); sp(pb, 0, 0)
run_el = pb.add_run()
br = OxmlElement('w:br'); br.set(qn('w:type'), 'page')
run_el._r.append(br)

# Answer key header
akt = doc.add_table(rows=1, cols=2); no_borders(akt)
rh(akt.rows[0], 1.0, 'exact')
for c in akt.rows[0].cells: shd(c, C_HEADER)

cw(akt.rows[0].cells[0], 14.5); va(akt.rows[0].cells[0])
p = akt.rows[0].cells[0].paragraphs[0]; sp(p)
r(p, "  ANSWER KEY  —  Present Simple Grammar", bold=True, size=14, color=WHITE)

cw(akt.rows[0].cells[1], 2.8); va(akt.rows[0].cells[1])
p = akt.rows[0].cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sp(p)
r(p, "A1–A2  ", bold=True, size=10, color=WHITE)

# Grammar reminder box
p = doc.add_paragraph(); sp(p, 6, 2)
r(p, "Grammar Reminder:", bold=True, size=11, color=C_EX_A)

tip_tbl = doc.add_table(rows=1, cols=1); tip_tbl.style = 'Table Grid'
shd(tip_tbl.rows[0].cells[0], "FFF9C4")
p = tip_tbl.rows[0].cells[0].paragraphs[0]; sp(p, 4, 2)
r(p, "Remember!  ", bold=True, size=10, color=C_HEADER)
r(p, "He/She/It → verb + -s/-es", size=10)
tip_lines = [
    ("(+)", "Subject + verb  →  He plays football."),
    ("(−)", "Subject + don't/doesn't + verb  →  She doesn't like coffee."),
    ("(?)", "Do/Does + subject + verb?  →  Does he play tennis?"),
    ("Signals:", "every day · always · usually · never · often · sometimes"),
]
for label, text in tip_lines:
    p = tip_tbl.rows[0].cells[0].add_paragraph(); sp(p, 2, 1)
    r(p, f"  {label}  ", bold=True, size=9.5, color=C_HEADER)
    r(p, text, size=9.5)

# Answer key — exercises
answers = {
    "A. Complete the table": [
        "Negative: don't / doesn't + play",
        "Question: Do / Does + subject + play?",
        "He/She/It: plays (affirmative)",
    ],
    "B. Circle correct option": [
        "1. plays   2. don't   3. Does   4. don't   5. watches   6. do"
    ],
    "C. Fill in the blanks": [
        "1. get up   2. doesn't like   3. go   4. has   5. watch"
    ],
    "D. Correct the mistakes": [
        "1. She doesn't like ice cream.",
        "2. Does he play tennis?",
        "3. They don't have a car.",
        "4. I go to bed at 10 pm.",
        "5. My brother plays the guitar.",
    ],
    "E. Write the questions": [
        "1. Do you like ice cream?",
        "2. Does she play tennis?",
        "3. Do they have a car?",
        "4. Does he go to bed at 10?",
    ],
}

for section, items in answers.items():
    p = doc.add_paragraph(); sp(p, 8, 2)
    r(p, section, bold=True, size=11, color=C_EX_A)
    for item in items:
        p = doc.add_paragraph(); sp(p, 1, 1)
        r(p, "  • ", bold=True, size=10, color=C_HEADER)
        r(p, item, size=10)

# Footer page 2
ft2 = doc.add_table(rows=1, cols=1); no_borders(ft2)
shd(ft2.rows[0].cells[0], C_FOOTER)
rh(ft2.rows[0], 0.48, 'exact')
p = ft2.rows[0].cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sp(p, 2, 2)
r(p, "workshido.com", bold=True, size=9, color=WHITE)
r(p, "  |  Page 2", size=8, color=WHITE)

# ── Kill phantom spacing ───────────────────────────────────────────────────────
kill_spacing(doc)

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
